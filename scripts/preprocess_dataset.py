"""模块三 5.1 数据预处理固定脚本（模块详细设计 5.1）。

用法: python preprocess_dataset.py <input> <out_dir> [--dataset-name NAME] [--task-type T]
                                 [--image-size WxH]

支持: CSV / Excel(.xlsx) / FASTA / PDB / 图片目录 / 压缩包(zip、tar)；
      旧版 .xls 识别后返回 unsupported（提示另存为 .xlsx）。

图片目录默认只做色彩空间归一（统一到 RGB，必要时保存归一副本），
传 --image-size 时额外缩放到目标尺寸（默认不缩放，避免破坏原始数据）。

输出: <out_dir>/preprocessed.csv（统一 schema）与
      <out_dir>/dataset.schema.json（统一 schema + alignment + 清洗统计）。
"""
import argparse
import csv
import hashlib
import json
import re
import shutil
import statistics
import sys
import tarfile
import zipfile
from collections import Counter
from pathlib import Path

UNIFIED_COLUMNS = ("id", "split", "label", "input")
META_PREFIX = "meta_"
# 标识/语义列：id 与 label 都不参与缺失值填充与 IQR 截断（口径见 clean 的 docstring）。
# 统一 schema 里这两列的名字由 UNIFIED_COLUMNS 固定，故此处取同一约定。
ID_COLUMN = "id"
LABEL_COLUMN = "label"

INPUT_ALIASES = ("input", "image", "img", "path", "file", "filename", "filepath",
                 "text", "sentence", "sequence", "seq")
LABEL_ALIASES = ("label", "labels", "class", "category", "target", "y", "tag")
ID_ALIASES = ("id", "index", "idx", "uid", "name")
SPLIT_ALIASES = ("split", "set", "phase", "subset")

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
TABLE_EXT = {".csv", ".tsv", ".xlsx", ".xls"}
UNSUPPORTED_EXT = {".fasta", ".fa", ".fna", ".pdb", ".ent"}
ARCHIVE_EXT = {".zip", ".tar", ".gz", ".tgz"}

# 图片目录中「划分」层级的目录名（不是类别名）；真实数据集的
# <root>/train/<class>/*.jpg 结构会把 train 误当标签，故需跳过并还原为 split。
SPLIT_DIR_NAMES = {
    "train": "train", "training": "train",
    "test": "test", "testing": "test",
    "val": "val", "valid": "val", "validation": "val", "dev": "val", "eval": "val",
}

MISSING = {"", "na", "n/a", "nan", "null", "none", "?", "-"}

# 单位统一（5.1 步骤 2 的量纲部分）：识别列名中的单位后缀并换算到统一量纲。
#
# 统一基准（括号内为换算后列名后缀）：
#   时间→s（ns/us/µs/ms/s/min/h…）、长度→m（nm/um/mm/cm/m/km）、
#   数据量→mb（字节，**按 1024 进制**：1 MB = 1024 KB = 1048576 B；不使用 1000 进制的
#   SI 口径，避免同一份数据因「1 KB 是 1000 还是 1024 字节」产生两种结果）、
#   比特→mbit（1 Mb = 1024 Kb，比特与字节分属不同量纲，绝不互相折算）、
#   比例→ratio（小数，% → 0.01、‰ → 0.001）、频率→hz、能量→j / wh、功率→w。
#
# 大小写敏感表优先命中：数据量里 Mb（兆比特）与 MB（兆字节）必须区分，
# 混同会把 1 Mb 当成 1 MB（差 8 倍）；查表顺序为「大小写敏感 → 小写回退」，
# 小写回退沿用既有口径（B/b 单字母都按字节处理；全小写 mb/gb 亦按字节）。
_UNIT_TABLE_CS = {
    "B": ("mb", 1 / 1048576), "b": ("mb", 1 / 1048576),
    "kB": ("mb", 1 / 1024), "KB": ("mb", 1 / 1024),
    "MB": ("mb", 1.0), "GB": ("mb", 1024.0), "TB": ("mb", 1048576.0),
    "KiB": ("mb", 1 / 1024), "MiB": ("mb", 1.0), "GiB": ("mb", 1024.0), "TiB": ("mb", 1048576.0),
    # 比特（小写 b 且带前缀）另立量纲 mbit
    "Kb": ("mbit", 1 / 1024), "Mb": ("mbit", 1.0),
    "Gb": ("mbit", 1024.0), "Tb": ("mbit", 1048576.0),
}

_UNIT_TABLE = {
    # 时间（基准 s）
    "ns": ("s", 1e-9), "us": ("s", 1e-6), "µs": ("s", 1e-6), "μs": ("s", 1e-6),
    "ms": ("s", 1e-3), "s": ("s", 1.0), "sec": ("s", 1.0), "secs": ("s", 1.0),
    "second": ("s", 1.0), "seconds": ("s", 1.0),
    "min": ("s", 60.0), "mins": ("s", 60.0), "minute": ("s", 60.0), "minutes": ("s", 60.0),
    "h": ("s", 3600.0), "hr": ("s", 3600.0), "hrs": ("s", 3600.0),
    "hour": ("s", 3600.0), "hours": ("s", 3600.0),
    # 长度（基准 m）
    "nm": ("m", 1e-9), "um": ("m", 1e-6), "µm": ("m", 1e-6), "μm": ("m", 1e-6),
    "mm": ("m", 1e-3), "cm": ("m", 1e-2), "dm": ("m", 1e-1), "m": ("m", 1.0),
    "km": ("m", 1000.0),
    # 数据量（字节，基准 mb，1024 进制；全小写写法按字节处理）
    "b": ("mb", 1 / 1048576), "kb": ("mb", 1 / 1024), "mb": ("mb", 1.0),
    "gb": ("mb", 1024.0), "tb": ("mb", 1048576.0),
    "kib": ("mb", 1 / 1024), "mib": ("mb", 1.0), "gib": ("mb", 1024.0), "tib": ("mb", 1048576.0),
    # 比例（基准 ratio，小数）
    "%": ("ratio", 0.01), "pct": ("ratio", 0.01), "percent": ("ratio", 0.01),
    "‰": ("ratio", 0.001), "permille": ("ratio", 0.001),
    # 频率（基准 hz）
    "hz": ("hz", 1.0), "khz": ("hz", 1e3), "mhz": ("hz", 1e6), "ghz": ("hz", 1e9),
    # 能量 / 功率（基准 j / wh / w；MJ/mJ、MW/mW 兆与毫同形，故意不收录以免误换算）
    "j": ("j", 1.0), "kj": ("j", 1e3),
    "wh": ("wh", 1.0), "kwh": ("wh", 1e3),
    "w": ("w", 1.0), "kw": ("w", 1e3),
}
# 要求单位前有分隔符（`latency_ms`、`duration(s)`），避免把 Mass/Address 这类词尾误判为单位
_UNIT_RE = re.compile(r"[_\s(（\-]+([A-Za-zµμ‰%]{1,8})[)）]?$")


def _unit_lookup(token: str) -> tuple[str, float] | None:
    """查单位表：先按原样（区分 Mb/MB、比特/字节），未命中再小写回退（沿用既有口径）。"""
    exact = _UNIT_TABLE_CS.get(token)
    if exact is not None:
        return exact
    return _UNIT_TABLE.get(token.lower())


def _fail(fmt: str, message: str) -> dict:
    return {"status": "unsupported", "format": fmt, "message": message}


def detect_format(path: Path) -> str:
    if path.is_dir():
        return "image_dir"
    ext = path.suffix.lower()
    if ext in UNSUPPORTED_EXT:
        return ext.lstrip(".")
    if ext in {".xlsx", ".xls"}:
        return "excel" if ext == ".xlsx" else "xls"
    if ext in {".csv", ".tsv"}:
        return "csv"
    if ext in IMAGE_EXT:
        return "image_dir"
    if ext in ARCHIVE_EXT or path.name.lower().endswith((".tar.gz", ".tar.bz2")):
        return "archive"
    return "unknown"


# ---------- 各格式读取：统一产出 (rows: list[dict], columns: list[str]) ----------

def _read_csv(path: Path) -> tuple[list[dict], list[str]]:
    for enc in ("utf-8-sig", "gbk", "latin-1"):
        try:
            with path.open("r", encoding=enc, newline="") as f:
                sample = f.read(8192)
                f.seek(0)
                delim = "\t" if path.suffix.lower() == ".tsv" else None
                if delim is None:
                    delim = "\t" if sample.count("\t") > sample.count(",") else ","
                reader = csv.DictReader(f, delimiter=delim)
                rows = [dict(r) for r in reader]
                return rows, list(reader.fieldnames or [])
        except UnicodeDecodeError:
            continue
    raise RuntimeError(f"无法解码 CSV 文件: {path}")


def _read_excel(path: Path) -> tuple[list[dict], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("缺少 openpyxl，无法读取 Excel") from e

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else f"col{i}" for i, c in enumerate(next(it))]
    rows = []
    for values in it:
        row = {header[i]: values[i] for i in range(min(len(header), len(values)))}
        if any(v is not None and str(v).strip() != "" for v in row.values()):
            rows.append(row)
    wb.close()
    return rows, header


def _read_fasta(path: Path) -> tuple[list[dict], list[str]]:
    """FASTA：``>id`` 起始一条记录，其后各行拼成序列（纯标准库，不引入 biopython）。"""
    rows: list[dict] = []
    current_id: str | None = None
    parts: list[str] = []
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        text = line.strip()
        if not text:
            continue
        if text.startswith(">"):
            if current_id is not None:
                rows.append({"id": current_id, "sequence": "".join(parts)})
            header = text[1:].split()
            current_id = (header[0] if header else "") or f"seq{len(rows) + 1}"
            parts = []
        elif current_id is not None:
            parts.append(text)
    if current_id is not None:
        rows.append({"id": current_id, "sequence": "".join(parts)})
    return rows, ["id", "sequence"]


def _read_pdb(path: Path) -> tuple[list[dict], list[str]]:
    """PDB：最少解析 ATOM/HETATM 固定列，产出每原子一行的表（纯标准库）。"""
    columns = ["id", "residue", "chain", "res_seq", "atom", "x", "y", "z"]
    rows: list[dict] = []
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        if not (line.startswith("ATOM") or line.startswith("HETATM")):
            continue
        try:
            x = float(line[30:38])
            y = float(line[38:46])
            z = float(line[46:54])
        except ValueError:
            continue
        atom = line[12:16].strip()
        chain = line[21:22].strip()
        res_seq = line[22:26].strip()
        rows.append({
            "id": f"{chain or '_'}:{res_seq}:{atom}",
            "residue": line[17:20].strip(),
            "chain": chain,
            "res_seq": res_seq,
            "atom": atom,
            "x": x,
            "y": y,
            "z": z,
        })
    return rows, columns


def _image_label_and_split(p: Path, root: Path) -> tuple[str, str]:
    """从图片所在目录推导 (label, split)。

    目录层级可能是 ``<root>/<split>/<class>/``、``<root>/<class>/<split>/``、
    只有 ``<root>/<class>/`` 或完全扁平（类别写在文件名里，如 ``dog.0.jpg``）：
    取沿途第一个「非 split 名」的目录作为 label，出现过的 split 名归一到 train/test/val；
    没有类别目录时退回文件名首个分隔符前的词，仍取不到则 unknown。
    """
    label, split = "", ""
    for part in p.parent.relative_to(root).parts:
        canonical = SPLIT_DIR_NAMES.get(part.casefold())
        if canonical:
            split = split or canonical
        elif not label:
            label = part
    if not label:
        label = re.split(r"[._\-\s]", p.stem)[0].strip() or "unknown"
    return label, split


def _read_image_dir(
    path: Path, out_dir: Path | None = None, target_size: tuple[int, int] | None = None
) -> tuple[list[dict], list[str], dict]:
    """读图片目录并做归一：色彩空间统一到 RGB，可选缩放到 target_size（默认不缩放）。

    需要归一（非 RGB 或指定了目标尺寸）时把副本写到 ``<out_dir>/_normalized/``，
    行内 path（即统一后的 input）指向该副本，使归一结果可持续访问。
    返回 (行, 列名, 归一统计)——统计落进 alignment.images。
    """
    try:
        from PIL import Image
    except ImportError as e:
        raise RuntimeError("缺少 Pillow，无法读取图片") from e

    rows: list[dict] = []
    stats = {
        "color_space": "RGB",
        "n_images": 0,
        "converted_to_rgb": 0,
        "target_size": list(target_size) if target_size else None,
        "resized": 0,
        "normalized_dir": None,
    }
    normalized_root: Path | None = None
    has_orig = False
    for p in sorted(path.rglob("*")):
        if p.suffix.lower() not in IMAGE_EXT:
            continue
        label, split = _image_label_and_split(p, path)
        try:
            with Image.open(p) as im:
                width, height, mode = im.width, im.height, im.mode
                orig_width, orig_height, orig_mode = width, height, mode
                needs_norm = mode != "RGB" or target_size is not None
                normed = False
                out_path = p
                if needs_norm and out_dir is not None:
                    if normalized_root is None:
                        normalized_root = out_dir / "_normalized"
                        stats["normalized_dir"] = str(normalized_root)
                    norm = im.convert("RGB")
                    if target_size is not None:
                        norm = norm.resize(target_size, getattr(Image, "Resampling", Image).LANCZOS)
                        stats["resized"] += 1
                    if mode != "RGB":
                        stats["converted_to_rgb"] += 1
                    out_path = normalized_root / p.relative_to(path)
                    out_path.parent.mkdir(parents=True, exist_ok=True)
                    norm.save(out_path)
                    # 行的 width/height/mode 必须描述**落盘的归一后文件**，否则下游按元数据
                    # 判断尺寸/通道会与真实文件不符；原图信息另存 orig_* 以便追溯
                    width, height, mode = norm.width, norm.height, norm.mode
                    normed = True
        except OSError:
            continue
        stats["n_images"] += 1
        row = {"path": str(out_path), "label": label, "width": width, "height": height, "mode": mode}
        if normed:
            row["orig_width"], row["orig_height"], row["orig_mode"] = orig_width, orig_height, orig_mode
            has_orig = True
        if split:
            row["split"] = split
        rows.append(row)
    columns = ["path", "label", "width", "height", "mode"]
    if any("split" in r for r in rows):
        columns.insert(1, "split")
    if has_orig:
        columns += ["orig_width", "orig_height", "orig_mode"]
    return rows, columns, stats


def _read_one(
    path: Path, out_dir: Path | None = None, target_size: tuple[int, int] | None = None
) -> tuple[list[dict], list[str], dict | None]:
    fmt = detect_format(path)
    if fmt == "csv":
        rows, columns = _read_csv(path)
        return rows, columns, None
    if fmt == "excel":
        rows, columns = _read_excel(path)
        return rows, columns, None
    if fmt in {"fasta", "fa", "fna"}:
        rows, columns = _read_fasta(path)
        return rows, columns, None
    if fmt in {"pdb", "ent"}:
        rows, columns = _read_pdb(path)
        return rows, columns, None
    if fmt == "image_dir":
        return _read_image_dir(path, out_dir, target_size)
    raise RuntimeError(f"不可直接读取的格式: {fmt}")


_MAX_ARCHIVE_MEMBERS = 200_000
_MAX_ARCHIVE_BYTES = 5 * 1024 ** 3  # 5 GiB，防解包炸弹


def _check_archive_limits(members: list) -> None:
    """解包前估算规模：条目数与解压后总字节超限即拒绝（含路径逃逸由 extractall(filter=data) 兜）。"""
    total = 0
    for _name, size in members:
        total += int(size or 0)
    if len(members) > _MAX_ARCHIVE_MEMBERS or total > _MAX_ARCHIVE_BYTES:
        raise RuntimeError(
            f"压缩包规模超限（条目 {len(members)}、解压后约 {total / 1024 ** 3:.1f} GiB），已拒绝解包"
        )


def _extract_archive(path: Path, dest: Path) -> None:
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            _check_archive_limits([(i.filename, i.file_size) for i in z.infolist()])
            z.extractall(dest)
        return
    if tarfile.is_tarfile(path):
        with tarfile.open(path) as t:
            _check_archive_limits([(m.name, m.size) for m in t.getmembers()])
            try:
                # filter="data" 由 stdlib 消毒：拒绝绝对路径、`..` 逃逸、符号链接与设备文件
                t.extractall(dest, filter="data")
            except TypeError:  # 极老 Python 无 filter 参数（当前环境 3.13，有）
                t.extractall(dest)
        return
    raise RuntimeError(f"无法识别的压缩包: {path}")


# ---------- 统一化 ----------

def _tokens(name: str) -> list[str]:
    return [t for t in re.split(r"[_\s(\-./]+", name.lower().strip()) if t]


def _match_alias(columns: list[str], aliases: tuple[str, ...]) -> str | None:
    """按「整名精确 → 分词精确」两级匹配列名。

    不做朴素子串匹配：别名里的 y、id 这类短词会命中 latency、valid 等无关列名。
    """
    lowered = {c.lower().strip(): c for c in columns}
    for alias in aliases:
        if alias in lowered:
            return lowered[alias]
    for col in columns:
        if any(token in aliases for token in _tokens(col)):
            return col
    return None


def _is_missing(v) -> bool:
    return v is None or str(v).strip().lower() in MISSING


def _to_float(v) -> float | None:
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


def _looks_like_sequence(values: list[str]) -> bool:
    """判定是否为序列列：绝大多数取值是纯字母，且典型长度达到序列量级。

    不要求每个值都够长——真实序列长度参差，短序列是正常的。
    """
    if not values:
        return False
    alpha_ratio = sum(1 for v in values if v.isalpha()) / len(values)
    lengths = sorted(len(v) for v in values)
    median_len = lengths[len(lengths) // 2]
    return alpha_ratio >= 0.8 and median_len >= 10


def unify(rows: list[dict], columns: list[str]) -> tuple[list[dict], dict, dict, list[dict]]:
    """映射到统一 schema；返回 (统一行, field_mapping, labels, warnings)。

    列名定位失败时**不静默**：定位不到输入列时统一行的 input 会整列为空字符串，
    此时在 warnings 里登记 {"code": "missing_input_column", ...}，并带上原始列名（columns）、
    可接受的别名（candidates，即 INPUT_ALIASES）与受影响行数，供人工复核与告警。
    warnings 是新增的第 4 项（追加在末尾）；调用方 `run()` 已配套修改，并把非空告警
    落进 dataset.schema.json 的 alignment.warnings。
    """
    col_in = _match_alias(columns, INPUT_ALIASES)
    col_label = _match_alias(columns, LABEL_ALIASES)
    col_id = _match_alias(columns, ID_ALIASES)
    col_split = _match_alias(columns, SPLIT_ALIASES)

    field_mapping: dict[str, str] = {}
    unified: list[dict] = []
    raw_labels: list[str] = []
    warnings: list[dict] = []
    if col_in is None:
        # 不静默：否则用户拿到的是 input 全空的数据集却看不出问题
        warnings.append({
            "code": "missing_input_column",
            "message": "未能从原始列中定位输入列（input），统一 schema 的 input 将全部为空字符串",
            "columns": [str(c) for c in columns],
            "candidates": list(INPUT_ALIASES),
            "affected_rows": len(rows),
        })

    for idx, row in enumerate(rows):
        out: dict = {}
        out["id"] = str(row.get(col_id)) if col_id and not _is_missing(row.get(col_id)) else str(idx)
        out["split"] = str(row.get(col_split)).strip() if col_split and not _is_missing(row.get(col_split)) else "train"
        out["label"] = str(row.get(col_label)).strip() if col_label and not _is_missing(row.get(col_label)) else ""
        out["input"] = str(row.get(col_in)).strip() if col_in and not _is_missing(row.get(col_in)) else ""
        if out["label"]:
            raw_labels.append(out["label"])
        mapped = {col_id, col_split, col_label, col_in}
        for c in columns:
            if c in mapped:
                continue
            out[META_PREFIX + c] = row.get(c)
        unified.append(out)

    for src, dst in ((col_id, "id"), (col_split, "split"), (col_label, "label"), (col_in, "input")):
        if src:
            field_mapping[src] = dst

    # 标签归并：同义写法（大小写/首尾空白）归并到出现最多的原始写法
    label_merge: dict[str, str] = {}
    if raw_labels:
        groups: dict[str, Counter] = {}
        for lab in raw_labels:
            groups.setdefault(lab.casefold(), Counter())[lab] += 1
        canonical = {k: v.most_common(1)[0][0] for k, v in groups.items()}
        for lab in dict.fromkeys(raw_labels):
            canon = canonical[lab.casefold()]
            if lab != canon:
                label_merge[lab] = canon
        for out in unified:
            if out["label"]:
                out["label"] = canonical[out["label"].casefold()]
    return unified, field_mapping, label_merge, warnings


def normalize_units(unified: list[dict], columns: list[str]) -> tuple[list[dict], list[str], dict]:
    """单位统一（5.1 步骤 2 的量纲部分）：列名带单位后缀的数值列换算到统一量纲并改名。

    返回 (统一行, 新列名列表, units 记录)。未命中任何单位时 units 为空。

    - **同量纲多单位合并**：`latency_ms` 与 `latency_us` 同归 `latency_s`，各列都换算到
      统一基准后合并进同一列（不再跳过）。每行取**列声明顺序**中第一个非空的来源列为准；
      同一行有多个来源列同时非空时按行计数 conflicts，如实记入 units。
      （注意：合并会抹平两列之间的差异，原本只靠这两列区分的行合并后可能变重复，
      由后续 `clean()` 的去重按既有口径处理。）
    - 目标列名与既有列（不参与本组合并者）重名时给出稳定消歧名 `<目标名>_2`、`_3`…，
      并在 units 记录里写明 disambiguation。
    - 保持既有行为：非数值列、非空值不足半数的列不换算（原因如实记 skipped，不静默）；
      meta_ 前缀约定与列改名规则不变。
    """
    units: dict = {}
    groups: dict[str, dict] = {}  # 目标列名 -> {target, base, sources[]}
    order: list[str] = []         # 目标列名的首次声明顺序

    for col in columns:
        if not col.startswith(META_PREFIX):
            continue
        orig = col[len(META_PREFIX):]
        m = _UNIT_RE.search(orig)
        if not m:
            continue
        token = m.group(1)
        found = _unit_lookup(token)
        if found is None:
            continue
        base, factor = found
        values = [_to_float(r.get(col)) for r in unified]
        present = [v for v in values if v is not None]
        # 仅对数值列换算：非空值须过半且都能转成数字（口径不变）；不换算的原因如实登记
        if not present or len(present) * 2 < len(unified):
            units[orig] = {"column": orig, "from": token, "to": base, "factor": factor,
                           "skipped": "非数值列或非空值不足半数，未换算"}
            continue
        target = f"{META_PREFIX}{orig[:m.start()].rstrip('_ -')}_{base}"
        group = groups.get(target)
        if group is None:
            group = {"target": target, "base": base, "sources": []}
            groups[target] = group
            order.append(target)
        group["sources"].append({"column": col, "orig": orig, "token": token, "factor": factor})

    # 目标列名与既有列重名（且该列不属于本组）→ 稳定消歧，避免产出重名列
    for target in list(order):
        group = groups[target]
        src_cols = {s["column"] for s in group["sources"]}
        if target not in columns or target in src_cols:
            continue
        candidate, n = target, 2
        while candidate in columns or candidate in groups:
            candidate = f"{target}_{n}"
            n += 1
        group["disambiguation"] = {"requested": target, "used": candidate,
                                   "reason": "目标列名与既有列重名，改用消歧名"}
        del groups[target]
        group["target"] = candidate
        groups[candidate] = group
        order[order.index(target)] = candidate

    col_target: dict[str, str] = {}
    for target in order:
        group = groups[target]
        sources = group["sources"]
        src_cols = [s["column"] for s in sources]
        merged_values: list = []
        conflicts = 0
        for row in unified:
            value = None
            non_null = 0
            for src in sources:
                num = _to_float(row.get(src["column"]))
                if num is None:
                    continue
                if value is None:
                    value = round(num * src["factor"], 10)  # 先声明者为准
                non_null += 1
            if non_null >= 2:
                conflicts += 1
            if value is None:
                # 全部来源都不是数值：沿用既有口径保留首个来源的原始值（缺失标记交给 clean 处理）
                value = row.get(src_cols[0])
            merged_values.append(value)
        for row, value in zip(unified, merged_values):
            row[target] = value
        for col in src_cols:
            col_target[col] = target
            if col != target:  # 同名（目标列本身就是来源列）时不能 pop，否则刚写入的值会被删掉
                for row in unified:
                    row.pop(col, None)
        for src in sources:
            record = {"column": target[len(META_PREFIX):], "from": src["token"],
                      "to": group["base"], "factor": src["factor"]}
            if len(sources) > 1:
                record["merged_from"] = [s["orig"] for s in sources]
                record["conflicts"] = conflicts
            if "disambiguation" in group:
                record["disambiguation"] = group["disambiguation"]
            units[src["orig"]] = record

    new_columns: list[str] = []
    emitted: set[str] = set()
    for col in columns:
        target = col_target.get(col)
        if target is None:
            new_columns.append(col)
        elif target not in emitted:
            new_columns.append(target)  # 目标列落在首个来源列的位置
            emitted.add(target)
    return unified, new_columns, units


def normalize_sequences(unified: list[dict]) -> tuple[list[dict], dict | None, int]:
    """序列规范（5.1 步骤 4）：统一大小写，并把超长序列截断到 IQR 上界。

    不补齐短序列（补齐会凭空造数据）；截断条数记入返回的第三项。
    """
    inputs = [str(r.get("input", "")).strip() for r in unified]
    present = [v for v in inputs if v]
    if not _looks_like_sequence(present):
        return unified, None, 0

    for row in unified:
        value = str(row.get("input", "")).strip()
        if value:
            row["input"] = value.upper()

    lengths = sorted(len(str(r["input"])) for r in unified if str(r.get("input", "")).strip())
    q1, q3 = _quantile(lengths, 0.25), _quantile(lengths, 0.75)
    upper = q3 + 1.5 * (q3 - q1)
    truncated = 0
    if upper > 0:
        for row in unified:
            value = str(row.get("input", ""))
            if len(value) > upper:
                row["input"] = value[: int(upper)]
                truncated += 1

    final = [len(str(r["input"])) for r in unified if str(r.get("input", "")).strip()]
    return unified, {
        "length_range": [min(final), max(final)],
        "case": "upper",
        "truncated": truncated,
    }, truncated


def clean(unified: list[dict], columns: list[str]) -> tuple[list[dict], dict]:
    """清洗：去重、缺失值、异常值。返回 (清洗后行, 统计)。

    口径（标识/语义列不参与数值修补，与既有 id 处理保持一致）：
    - **id 与 label 不参与缺失值填充，也不参与 IQR 异常值截断。** 对 id 补值/裁剪会伪造或重复
      主键（数值型 id 被填成中位数即重复）；label 是类别语义，填众数等于凭空造类别
      （把缺失样本划进多数类），IQR 截断会改写类别取值（数值型 label 被压到分位边界）。
    - 整行缺失比例判据**沿用既有口径不变**（分母为除 id 外的列、含 label，>50% 才删行）：
      label 缺失的行按该判据决定去留，保留下来就保持为空，绝不用众数伪造。
    - 去重的「内容字段」也**沿用既有口径不变**：除 id 外的全部字段（含 label、split 与 meta_*）。
    - 统计口径：`excluded_from_fill_and_clip` 列出未参与填充/截断的列；`label_missing` 如实记录
      清洗后 label 仍为空的行数；`filled` / `outliers_clipped` 里不会出现这两列，
      可据此复核「label 没有被众数填充、也没有被 IQR 截断」。
    """
    stats = {
        "duplicates_removed": 0,
        "rows_dropped_missing": 0,
        "filled": {},
        "outliers_clipped": {},
        "excluded_from_fill_and_clip": [],
        "label_missing": 0,
    }

    # 去重：以内容字段的哈希为键（排除生成的 id，否则行号不同会导致重复行无法识别）
    # 注意：内容字段含 label 与 meta_*，只有 meta 列不同的行哈希不同，不会被误判为重复
    seen: set[str] = set()
    deduped: list[dict] = []
    for row in unified:
        content = {k: v for k, v in row.items() if k != ID_COLUMN}
        key = hashlib.sha1(json.dumps(content, sort_keys=True, default=str).encode("utf-8")).hexdigest()
        if key in seen:
            stats["duplicates_removed"] += 1
            continue
        seen.add(key)
        deduped.append(row)

    # 不参与填充/截断的标识列（只登记实际存在的列，便于复核基数）
    stats["excluded_from_fill_and_clip"] = [c for c in (ID_COLUMN, LABEL_COLUMN) if c in columns]

    # 缺失值：整行缺失 >50% 删除；数值列填中位数、其余填众数
    # id 列不参与填充/截断——对它补值或裁剪会伪造/重复主键（数值型 id 被填成中位数即重复）；
    # label 列同理不参与——填众数会伪造类别、IQR 截断会改写类别边界。
    # 整行缺失比例沿用既有口径（分母仍为除 id 外的列，含 label），不因新增排除而改变判据。
    value_cols = [c for c in columns if c != ID_COLUMN]
    fillable_cols = [c for c in value_cols if c != LABEL_COLUMN]
    kept: list[dict] = []
    for row in deduped:
        missing = sum(1 for c in value_cols if _is_missing(row.get(c)))
        if value_cols and missing / len(value_cols) > 0.5:
            stats["rows_dropped_missing"] += 1
            continue
        kept.append(row)

    numeric_cols: list[str] = []
    for c in fillable_cols:
        vals = [row.get(c) for row in kept if not _is_missing(row.get(c))]
        if vals and all(_to_float(v) is not None for v in vals):
            numeric_cols.append(c)

    for c in fillable_cols:
        present = [row.get(c) for row in kept if not _is_missing(row.get(c))]
        if not present:
            continue
        if c in numeric_cols:
            fill = statistics.median([_to_float(v) for v in present])
        else:
            fill = Counter(str(v) for v in present).most_common(1)[0][0]
        filled = 0
        for row in kept:
            if _is_missing(row.get(c)):
                row[c] = fill
                filled += 1
        if filled:
            stats["filled"][c] = filled

    # 异常值：IQR 1.5 倍，超出者截断到边界并计数（numeric_cols 已排除 id/label）
    for c in numeric_cols:
        vals = sorted(_to_float(row.get(c)) for row in kept)
        if len(vals) < 4:
            continue
        q1 = _quantile(vals, 0.25)
        q3 = _quantile(vals, 0.75)
        iqr = q3 - q1
        lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        clipped = 0
        for row in kept:
            v = _to_float(row.get(c))
            if v is None:
                continue
            if v < lo:
                row[c] = lo
                clipped += 1
            elif v > hi:
                row[c] = hi
                clipped += 1
        if clipped:
            stats["outliers_clipped"][c] = clipped
    # 如实计数：清洗后 label 仍为空的行数（这些行没有被众数填充，保留为空）；
    # 列集合里没有 label 列时不计（否则会把「无标签数据集」误算成「全部标签缺失」）
    if LABEL_COLUMN in columns:
        stats["label_missing"] = sum(1 for row in kept if _is_missing(row.get(LABEL_COLUMN)))
    return kept, stats


def _quantile(sorted_vals: list[float], q: float) -> float:
    if not sorted_vals:
        return 0.0
    pos = (len(sorted_vals) - 1) * q
    lo, hi = int(pos), min(int(pos) + 1, len(sorted_vals) - 1)
    return sorted_vals[lo] + (sorted_vals[hi] - sorted_vals[lo]) * (pos - lo)


# ---------- 入口 ----------

def _write_outputs(out_dir: Path, rows: list[dict], columns: list[str]) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_path = out_dir / "preprocessed.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            writer.writerow({c: row.get(c) for c in columns})
    return csv_path


def run(
    input_path: Path,
    out_dir: Path,
    dataset_name: str,
    task_type: str,
    image_size: tuple[int, int] | None = None,
) -> dict:
    fmt = detect_format(input_path)
    if fmt == "xls":
        return _fail(fmt, "暂不支持旧版 .xls（请另存为 .xlsx）")
    if fmt == "unknown":
        return _fail(fmt, f"无法识别格式: {input_path.name}")

    image_stats: dict | None = None
    if fmt == "archive":
        # 解压到产出目录下的 _files/：图片类数据集的统一 CSV 里 input 必须指向可持续访问的路径，
        # 若解压到临时目录，本次运行结束即被清理，CSV 中的图片路径会全部失效（真实数据集暴露的问题）。
        extract_dir = out_dir / "_files"
        if extract_dir.exists():
            shutil.rmtree(extract_dir, ignore_errors=True)
        extract_dir.mkdir(parents=True, exist_ok=True)
        _extract_archive(input_path, extract_dir)
        inner = sorted(p for p in extract_dir.rglob("*") if p.is_file())
        if not inner:
            return _fail(fmt, "压缩包内没有文件")
        # 递归一次：优先取表格，其次取图片/子目录
        table = next((p for p in inner if p.suffix.lower() in TABLE_EXT), None)
        if table is not None:
            rows, columns, image_stats = _read_one(table, out_dir, image_size)
            fmt = detect_format(table)
        else:
            img_files = [p for p in inner if p.suffix.lower() in IMAGE_EXT]
            if not img_files:
                return _fail(fmt, "压缩包内没有可识别的表格或图片")
            # 仅当压缩包只有一个顶层目录（常见的「多包一层」导出）才下钻，
            # 否则以解压根目录为根：train/、test/ 并列时只取第一个会把一半数据丢掉。
            tops = {p.relative_to(extract_dir).parts[0] for p in img_files}
            single_top = next(iter(tops)) if len(tops) == 1 else None
            img_root = extract_dir
            if single_top and (extract_dir / single_top).is_dir():
                img_root = extract_dir / single_top
            rows, columns, image_stats = _read_image_dir(img_root, out_dir, image_size)
            fmt = "image_dir"
    else:
        rows, columns, image_stats = _read_one(input_path, out_dir, image_size)

    if not rows:
        return _fail(fmt, "未读取到任何数据行")

    unified, field_mapping, label_merge, warnings = unify(rows, columns)
    unified_columns = list(UNIFIED_COLUMNS) + [c for c in unified[0] if c.startswith(META_PREFIX)]
    # 顺序：单位换算 → 序列规范 → 清洗（清洗在统一量纲后进行，异常值判定才准确）
    unified, unified_columns, units = normalize_units(unified, unified_columns)
    unified, sequence, seq_truncated = normalize_sequences(unified)
    cleaned, stats = clean(unified, unified_columns)
    if seq_truncated:
        stats["sequences_truncated"] = seq_truncated

    csv_path = _write_outputs(out_dir, cleaned, unified_columns)
    labels = sorted({r["label"] for r in cleaned if r.get("label")})
    summary = {
        "status": "ok",
        "format": fmt,
        "dataset_name": dataset_name,
        "task_type": task_type,
        "n_rows": len(cleaned),
        "n_columns": len(unified_columns),
        "fields": unified_columns,
        "labels": labels,
        "cleaning": stats,
        "output_csv": str(csv_path),
    }
    summary["alignment"] = {
        "field_mapping": field_mapping,
        "label_merge": label_merge,
        "sequence": sequence,
        "units": units,
        "version": "1.0",
    }
    if warnings:
        # 统一化告警（如未定位到输入列）随 alignment 落进 dataset.schema.json，不静默丢弃
        summary["alignment"]["warnings"] = warnings
    if image_stats is not None:
        # 图片归一统计（色彩空间/RGB 转换/缩放）落进 alignment.images（5.6 约定结构）
        summary["alignment"]["images"] = image_stats
    (out_dir / "dataset.schema.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return summary


def main() -> int:
    ap = argparse.ArgumentParser(description="模块三 5.1 数据预处理")
    ap.add_argument("input")
    ap.add_argument("out_dir")
    ap.add_argument("--dataset-name", default="")
    ap.add_argument("--task-type", default="classification")
    ap.add_argument("--image-size", default="", help="图片目标尺寸 WxH（如 224x224），留空则不缩放")
    args = ap.parse_args()

    image_size: tuple[int, int] | None = None
    if args.image_size.strip():
        m = re.match(r"^(\d+)\s*[xX*×]\s*(\d+)$", args.image_size.strip())
        if not m:
            print(json.dumps(
                {"status": "error", "message": "--image-size 需为 WxH 形式，如 224x224"}, ensure_ascii=False
            ))
            return 1
        image_size = (int(m.group(1)), int(m.group(2)))

    input_path = Path(args.input)
    if not input_path.exists():
        print(json.dumps({"status": "error", "message": f"输入不存在: {input_path}"}, ensure_ascii=False))
        return 1
    try:
        summary = run(input_path, Path(args.out_dir), args.dataset_name, args.task_type, image_size)
    except Exception as e:  # noqa: BLE001
        print(json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False))
        return 1
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
